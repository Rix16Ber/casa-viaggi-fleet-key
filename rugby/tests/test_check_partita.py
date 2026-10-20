import json
import tempfile
import unittest
from pathlib import Path

import openpyxl

from rugby.check_partita import main

FIXTURE = Path(__file__).resolve().parent / "fixtures" / "conteggi-esempio.csv"
HEADERS = [
    "N",
    "Inizio",
    "Fine",
    "Da dove",
    "Dove",
    "Start set piece",
    "End set piece",
    "Nota Veo",
    "Metri",
    "Esito",
]


def _fase(number, da, dove, metri, esito="Positiva"):
    return {
        "N": number,
        "Inizio": "00:10",
        "Fine": "01:00",
        "Da dove": da,
        "Dove": dove,
        "Start set piece": "Touche",
        "End set piece": "Touche",
        "Nota Veo": "fase",
        "Metri": metri,
        "Esito": esito,
    }


def _workbook(path: Path, phases: list[dict]) -> None:
    book = openpyxl.Workbook()
    piece = book.active
    piece.title = "Set piece"
    row = 5
    for index in range(7):
        piece.cell(row, 1, "Mischia")
        piece.cell(row, 2, f"00:{index:02d}")
        piece.cell(row, 3, "Monte")
        piece.cell(row, 4, "Vinta")
        piece.cell(row, 6, "Vinta")
        row += 1
    for index in range(11):
        piece.cell(row, 1, "Mischia")
        piece.cell(row, 2, f"10:{index:02d}")
        piece.cell(row, 3, "Padova")
        piece.cell(row, 4, "Vinta")
        piece.cell(row, 6, f"Reset {index}:00" if index < 9 else "Vinta")
        row += 1
    for index in range(27):
        piece.cell(row, 1, "Touche")
        piece.cell(row, 2, f"20:{index:02d}")
        piece.cell(row, 3, "Monte")
        piece.cell(row, 4, "Vinta")
        piece.cell(row, 6, "Vinta")
        row += 1
    events = book.create_sheet("Eventi")
    for index in range(5):
        events.cell(4 + index, 2, f"30:{index:02d}")
        events.cell(4 + index, 3, "Monte")
        events.cell(4 + index, 4, "Meta")
    for index in range(2):
        events.cell(9 + index, 2, f"40:{index:02d}")
        events.cell(9 + index, 3, "Padova")
        events.cell(9 + index, 4, "Trasformazione")
    phases_sheet = book.create_sheet("Fasi Gioco")
    for column, header in enumerate(HEADERS, start=1):
        phases_sheet.cell(1, column, header)
    for offset, phase in enumerate(phases):
        for column, header in enumerate(HEADERS, start=1):
            phases_sheet.cell(2 + offset, column, phase[header])
    book.save(path)
    book.close()


class CheckPartitaTest(unittest.TestCase):
    def test_pass_si_rilancia_senza_scrivere_il_file(self):
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw)
            workbook = root / "partita.xlsx"
            _workbook(workbook, [_fase(1, "50", "22 CUS", "Positiva")])
            before = workbook.read_bytes()
            log = root / "log.txt"
            first = main([str(workbook), str(FIXTURE)], log)
            second = main([str(workbook), str(FIXTURE)], log)
            self.assertEqual(first, 0)
            self.assertEqual(second, 0)
            self.assertEqual(workbook.read_bytes(), before)
            rows = [json.loads(line) for line in log.read_text(encoding="utf-8").splitlines()]
            self.assertEqual([row["status"] for row in rows], ["ok", "ok"])
            self.assertEqual(rows[0]["fail"], 0)

    def test_fail_esce_uno(self):
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw)
            workbook = root / "vuoto.xlsx"
            book = openpyxl.Workbook()
            book.active.title = "Set piece"
            book.create_sheet("Eventi")
            phases = book.create_sheet("Fasi Gioco")
            for column, header in enumerate(HEADERS, start=1):
                phases.cell(1, column, header)
            book.save(workbook)
            book.close()
            code = main([str(workbook), str(FIXTURE)], root / "log.txt")
            self.assertEqual(code, 1)
            record = json.loads((root / "log.txt").read_text(encoding="utf-8"))
            self.assertEqual(record["status"], "errore")
            self.assertGreater(record["fail"], 0)

    def test_warn_non_e_un_fail(self):
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw)
            workbook = root / "incrociate.xlsx"
            _workbook(
                workbook,
                [
                    _fase(1, "50", "22 CUS", "Positiva"),
                    _fase(2, "22", "10", "Positiva"),
                ],
            )
            code = main([str(workbook), str(FIXTURE)], root / "log.txt")
            self.assertEqual(code, 0)
            record = json.loads((root / "log.txt").read_text(encoding="utf-8"))
            self.assertEqual(record["warn"], 1)
            self.assertEqual(record["fail"], 0)


class SorgenteSolaLetturaTest(unittest.TestCase):
    def test_niente_save_sullo_xlsx(self):
        source = Path(__file__).resolve().parents[1] / "check_partita.py"
        self.assertNotIn(".save(", source.read_text(encoding="utf-8"))
