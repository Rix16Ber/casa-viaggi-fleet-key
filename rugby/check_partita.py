#!/usr/bin/env python3
"""Controllo file partita (riscaricato da OneDrive). Uso: check-partita.py file.xlsx veo-conteggi.csv [badge.csv]
badge.csv opzionale: tipo,orario,squadra (PMA/CUS) dal badge Veo."""
import json
import sys
import csv
import re
import openpyxl
from datetime import datetime, timezone
from pathlib import Path

_STEP = "pre-consegna"


def main(argv=None, log_path=None) -> int:
    if argv is None:
        argv = sys.argv[1:]
    saved = sys.argv
    sys.argv = ["check-partita.py", *argv]
    try:
        return _controlla(Path(log_path) if log_path else Path("cursor-rugby.log"))
    finally:
        sys.argv = saved


def _controlla(log_path: Path) -> int:
    f,veo=sys.argv[1],sys.argv[2]; badgef=sys.argv[3] if len(sys.argv)>3 else None
    wb=openpyxl.load_workbook(f,data_only=True)
    res=[]
    def rep(stato,nome,det=''): res.append((stato,nome,det)); print(f'[{stato}] {nome}'+(f'\n      '+det.replace('\n','\n      ') if det else ''))
    T={'Monte':'PMA','Padova':'CUS'}
    # ---- Set piece
    sp=wb['Set piece']; rows=[]
    for r in range(5,sp.max_row+1):
        a=sp.cell(r,1).value
        if a in ('Touche','Mischia'): rows.append(dict(r=r,tipo=a,ora=sp.cell(r,2).value,chi=sp.cell(r,3).value,esito=sp.cell(r,4).value,nota=sp.cell(r,6).value))
    V={x['tipo']:x for x in csv.DictReader(open(veo))}
    def n(x): return int(x) if x not in (None,'') else None
    print('== 1) Totali file vs Veo')
    M=[x for x in rows if x['tipo']=='Mischia']; resets={'PMA':0,'CUS':0}; reset_list=[]
    for x in M:
        k=len(re.findall(r'Reset \d+:\d+',x['nota'] or '')); resets[T.get(x['chi'],'?')]=resets.get(T.get(x['chi'],'?'),0)+k
        reset_list+= [(t,x['ora']) for t in re.findall(r'Reset (\d+:\d+)',x['nota'] or '')]
    v=V.get('Scrum')
    if v:
        righe={s:sum(1 for x in M if T.get(x['chi'])==s) for s in ('PMA','CUS')}
        tot=len(M)+sum(resets.values())
        det=f"righe file {len(M)} (PMA {righe['PMA']}, CUS {righe['CUS']}) + reset in nota {sum(resets.values())} (PMA {resets['PMA']}, CUS {resets['CUS']}) = {tot} (PMA {righe['PMA']+resets['PMA']}, CUS {righe['CUS']+resets['CUS']}); Veo {v['totale']} (PMA {v['PMA']}, CUS {v['CUS']}, reset {v['reset']})"
        ok= tot==n(v['totale']) and (not v['PMA'] or righe['PMA']+resets['PMA']==n(v['PMA'])) and (not v['CUS'] or righe['CUS']+resets['CUS']==n(v['CUS'])) and (not v['reset'] or sum(resets.values())==n(v['reset']))
        rep('PASS' if ok else 'FAIL','Mischie',det)
    L=[x for x in rows if x['tipo']=='Touche']; v=V.get('Lineout')
    if v:
        det=f"file {len(L)} (PMA {sum(T.get(x['chi'])=='PMA' for x in L)}, CUS {sum(T.get(x['chi'])=='CUS' for x in L)}); Veo {v['totale']}; differenza {n(v['totale'])-len(L)}"
        rep('PASS' if len(L)==n(v['totale']) else 'FAIL','Touche',det+("\nIl file non dice quali touche Veo mancano: servono gli orari Veo per il confronto riga per riga." if len(L)!=n(v['totale']) else ''))
    ev=wb['Eventi']; E=[(ev.cell(r,2).value,ev.cell(r,3).value,ev.cell(r,4).value) for r in range(4,ev.max_row+1) if ev.cell(r,4).value]
    for tipo,lab in (('Try','Meta'),('Conversion','Trasformazione')):
        v=V.get(tipo)
        if not v: continue
        xs=[e for e in E if e[2]==lab]
        rif=n(v.get('riusciti')) if v.get('riusciti') else n(v['totale'])
        cosa='riuscite su Veo' if v.get('riusciti') else 'Veo'
        det=f"file {len(xs)} ("+', '.join(f'{o} {T.get(s,s)}' for o,s,_ in xs)+f"); {cosa} {rif}"+(f" (tentativi Veo {v['totale']})" if v.get('riusciti') else '')+f"; differenza {rif-len(xs)}"
        rep('PASS' if len(xs)==rif else 'FAIL',f'{lab} (Veo {tipo})',det)
    print('== 2) Celle vuote')
    vu=[f"riga {x['r']} {x['tipo']} {x['ora']}: "+'/'.join(k for k in ('chi','nota') if not x[k]) for x in rows if not x['chi'] or not x['nota']]
    rep('FAIL' if vu else 'PASS','Set piece Squadra/Nota Veo','\n'.join(vu))
    da=[f"{x['tipo']} {x['ora']} ({x['chi']})" for x in rows if not x['esito']]
    rep('INFO',f'Set piece Esito da compilare a Riccardo: {len(da)}','\n'.join(da))
    fg=wb['Fasi Gioco']; H=[c.value for c in fg[1]]; ci={h:i for i,h in enumerate(H)}
    F=[]
    for r in range(2,fg.max_row+1):
        row=[c.value for c in fg[r]]
        if row[ci['N']] is None: continue
        F.append(dict(zip(H,row)))
    vu=[f"fase {x['N']} ({x['Inizio']}): "+'/'.join(k for k in ('Nota Veo','Metri','Esito') if x.get(k) in (None,'')) for x in F if any(x.get(k) in (None,'') for k in ('Nota Veo','Metri','Esito'))]
    rep('FAIL' if vu else 'PASS',f'Fasi Gioco Nota/Metri/Esito ({len(F)} fasi)','\n'.join(vu))
    print('== 3) Regole')
    def sec(t):
        a=[int(v) for v in str(t).split(':')]; return a[0]*60+a[1]
    mete=[(sec(o),T.get(sq,sq)) for o,sq,ti in E if ti=='Meta' and o]
    def pos(line,x,inizio):
        """distanza dalla linea di meta CUS; linee senza squadra = nostra metà (10 = 10 m dalla metà campo)"""
        s=(line or '').strip().upper().replace(' ','')
        if s.startswith('META'):
            if 'CUS' in s: return 0
            if inizio: return 100 if 'LIBERAZIONE' in str(x.get('Start set piece')).upper() else 0
            try:
                chi=[t for o,t in mete if sec(x['Inizio'])<=o<=sec(x['Fine'])+5]
            except Exception: chi=[]
            if chi: return 100 if chi[0]=='CUS' else 0
            txt=(str(x.get('Start set piece'))+' '+str(x.get('End set piece'))).upper()
            return 100 if 'CUS' in txt else 0
        m=re.match(r'(\d+)(CUS|MONTE|PMA)?$',s)
        if not m: return None
        d=int(m.group(1)); cus=m.group(2)=='CUS'
        if d==10: return 40 if cus else 60
        if d==50: return 50
        return d if cus else 100-d
    nostre=[f"fase {x['N']}: {x['Da dove']} → {x['Dove']}" for x in F if any(re.fullmatch(r'\s*\d+\s*',str(x.get(k) or '')) for k in ('Da dove','Dove'))]
    incr=[]
    for a,b in zip(F,F[1:]):
        da,db=str(a.get('Dove') or ''),str(b.get('Da dove') or '')
        if re.sub(r'\D','',da)==re.sub(r'\D','',db) and da and db and ('CUS' in da)!=('CUS' in db) and re.sub(r'\D','',da)!='50':
            incr.append(f"fase {a['N']} finisce «{da}» / fase {b['N']} parte «{db}»")
    rep('INFO',f'Linee senza squadra lette come nostre: {len(nostre)} fasi')
    rep('WARN' if incr else 'PASS','Linee incrociate con fasi vicine (stessa linea, squadra diversa)','\n'.join(incr))
    bad=[f"fase {x['N']} {x['Inizio']}: «{x['End set piece']}» Esito={x['Esito']}" for x in F if re.match(r'calcio(?!.*contro)',str(x.get('End set piece') or ''),re.I) and 'CUS' not in str(x.get('End set piece')).upper() and str(x.get('Esito')).lower()!='positiva']
    rep('FAIL' if bad else 'PASS','Calci senza CUS = nostri e Esito Positivo','\n'.join(bad))
    sc=[]
    for x in rows:
        d=str(x['esito'] or '').lower()
        if d=='calcio contro' or d=='calcio favore' or d=='calcio a favore':
            ec=sp.cell(x['r'],5).value; att='persa' if 'contro' in d else 'vinta'
            if ec!=att: sc.append(f"{x['tipo']} {x['ora']}: D={x['esito']} conteggio={ec} atteso {att}")
    for x in F:
        e=str(x.get('End set piece') or '').lower()
        if re.match(r'calcio\s*contro',e) and str(x['Esito']).lower()!='negativa': sc.append(f"fase {x['N']}: «{x['End set piece']}» Esito={x['Esito']} atteso Negativa")
    rep('FAIL' if sc else 'PASS','Calcio contro → Negativa/persa, calcio favore → Positiva/vinta','\n'.join(sc))
    if badgef:
        B={(b['tipo'],b['orario']):b['squadra'] for b in csv.DictReader(open(badgef))}
        bb=[f"{x['tipo']} {x['ora']}: file {x['chi']} / badge {B[(x['tipo'],x['ora'])]}" for x in rows if (x['tipo'],x['ora']) in B and T.get(x['chi'])!=B[(x['tipo'],x['ora'])]]
        miss=[f"{x['tipo']} {x['ora']}" for x in rows if (x['tipo'],x['ora']) not in B]
        rep('FAIL' if bb else 'PASS',f'Squadra set piece = badge Veo ({len(B)} badge)','\n'.join(bb+(['senza badge: '+', '.join(miss)] if miss else [])))
    mm=[]
    for x in F:
        p0,p1=pos(x.get('Da dove'),x,True),pos(x.get('Dove'),x,False)
        if p0 is None or p1 is None: mm.append(f"fase {x['N']}: non leggibile «{x.get('Da dove')} → {x.get('Dove')}» Metri={x['Metri']}"); continue
        att='Positiva' if p1<p0 else 'Negativa' if p1>p0 else 'Medio'
        if str(x['Metri']).lower().rstrip('a').rstrip('o')!=att.lower().rstrip('a').rstrip('o'):
            mm.append(f"fase {x['N']}: {x.get('Da dove')} → {x.get('Dove')} atteso {att}, file {x['Metri']}")
    rep('FAIL' if mm else 'PASS','Metri coerenti con Da dove → Dove','\n'.join(mm))
    fails = sum(s == "FAIL" for s, _, _ in res)
    warns = sum(s == "WARN" for s, _, _ in res)
    print('== ESITO:', 'FAIL' if fails else 'PASS', f"({fails} FAIL, {warns} WARN)")
    _log(log_path, "errore" if fails else "ok", fail=fails, warn=warns, file=str(f))
    return 1 if fails else 0


def _log(path: Path, status: str, **fields: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    record = {
        "ts": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "step": _STEP,
        "status": status,
    }
    record.update(fields)
    with path.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(record, ensure_ascii=False) + "\n")


if __name__ == "__main__":
    sys.exit(main())
